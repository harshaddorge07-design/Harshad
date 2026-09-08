#=========================================================
#Import
#=========================================================
import openpyxl
import uuid
import os 
import time
import datetime
from rich.console import Console
from pyfiglet import figlet_format
from pwinput import pwinput
from rich.progress import track
console = Console()

#=========================================================
# Check balance
#=========================================================
def check_bal(acc_no):
    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active

    for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
            if (row[0] == acc_no):
                account= row[7]
                break

    wb.close()

#=========================================================
# Sub menu 
#=========================================================
def sub_menu(acc_no):
    while True:
        console.print("""
        1. Check balance
        2. Deposite money
        3. Withdraw money
        4. Transition history
        5.Logout
        """)
        choice = input("Enter yor choice :: ")
        if choice == "1":
            check_bal(acc_no)
        elif choice == "2":
            depo_money(acc_no)
        elif choice == "3":
            pass
        elif choice == "4":
            pass
        elif choice == "5":
            pass
        else:
            err_msg("Choice")

#=========================================================
# Login
#=========================================================
def login():
    console.print("[bold yellow]LOGIN [/bold yellow]")
    acc_no = input("Enter your account no. :: ")
    pin = pwinput("Enter your pin :: ", mask="*")
    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active
    for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        if (row[0] == acc_no) and (row[-1] == pin):
            user_name = row[1]
            break
    wb.close()
    console.print("[green]Verfying account ....[/green]")
    progress_bar(5)
    console.print("[bold green]LOGIN SUCCESSFUL[/bold green]")
    console.print(f"Welcome [cyan]{user_name}[/cyan] ")
    print()
    sub_menu(acc_no)
    console.print("[bold magenta]Thank you visit again[/bold magenta]")


#=========================================================
# current day and time
#========================================================= 
def current_dt(): #dt = date & time 
    now = datetime.datetime.now()

    formatted_1 = now.strftime("%Y-%m-%d %H:%M:%S")
    return formatted_1
    # print(formatted_1) # Output: 2026-09-08 09:30:00

#=========================================================
# Process bar 
#=========================================================
def progress_bar(val=1):
    for step in track(range(val), description="Processing..."):
        time.sleep(1)

#=========================================================
#INITIALIZATION
#=========================================================
file_name = "bank_records.xlsx"
bank_name = figlet_format("BY BANK", font="slant")

#=========================================================
# error message 
#=========================================================
def err_msg(error):
    console.print(f"[bold red]Invalid {error}, please try again [/bold red]")

#=========================================================
# Excel Creation Function
#=========================================================
def create_excel_file():
    """this function create a excel name bank_records.xlsx"""
    if not os.path.exists(file_name):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Bank Records"
        headers = ["Account No", "Name", "PIN", "Transaction ID", "transaction Type", "Amount", 
                   "Previous Balance", "Current Balance", "Date-Time"]
        sheet.append(headers)

        wb.save(file_name)
        wb.close()
        console.print("[bold cyan]Excel File created successfully[/bold cyan]")

#=========================================================
# Create Account 
#=========================================================
def create_account():
    console.print("[bold yellow]CREATE ACCOUNT[/Bold yellow]")
    name = input("Enter user name :: ")
    acc_no = input("Enter account number :: ")
    if not acc_no.isdigit():
        err_msg("Account number")
        return
    if acc_no == "exists":
        pass
    pin = pwinput("Enter your 4 digit PIN :: ")
    if len(pin) != 4 or not pin.isdigit():
        err_msg("PIN")
        return

    try:
        opening_bal = float(input("Enter a opening balance :: "))
        if opening_bal < 0:
            err_msg("Opening balance")
            return
    except ValueError:
        err_msg("Value Entered")

    transition_id = str(uuid.uuid4())[:8]
    status = "Opening"
    dt  = current_dt()

    wb = openpyxl.load_workbook(file_name)
    sheet = wb.active
    sheet.append([acc_no, name, pin, transition_id, status, opening_bal,0 , opening_bal, dt ])
    wb.save(file_name)
    wb.close()

    progress_bar(5)

    console.print("[bold yellow] Account Created successfully[/bold yellow]")

#=========================================================   
#Main 
# #=========================================================     
if __name__ == "__main__":
    create_excel_file()
    while True:
        console.print(f"[green]{bank_name}[/green]")
        print()
        console.print("[bold cyan]Welcome to Python Banking System[/bold cyan]")
        print("""1.Create Account
2.Login
3.Exit""")
        print()
        choice = console.input("[bold purple]Enter your choice :: [/bold purple]")
        if choice == "1":
            create_account()
        elif choice == "2":
            login()   
        elif choice == "3":
            break

os.startfile(file_name)

